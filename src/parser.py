from pathlib import Path

import pymupdf
from docx import Document


def extract_pdf_text(path: Path) -> str:
    """Extract text from a PDF."""
    with pymupdf.open(path) as doc:
        pages = [
            page.get_text("text")
            for page in doc
        ]

    return "\n".join(pages).strip()


def extract_docx_text(path: Path) -> str:
    """Extract text from a DOCX file, including tables."""
    document = Document(path)

    parts = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Extract table content
    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                text = cell.text.strip()

                if text:
                    cells.append(text)

            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def extract_txt_text(path: Path) -> str:
    """Extract text from a TXT file."""
    return path.read_text(
        encoding="utf-8",
        errors="replace",
    ).strip()


def extract_resume_text(path: Path) -> str:
    """
    Automatically select the appropriate parser
    based on the resume file extension.
    """

    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(path)

    if extension == ".docx":
        return extract_docx_text(path)

    if extension == ".txt":
        return extract_txt_text(path)

    raise ValueError(
        f"Unsupported file format: {path.suffix}"
    )